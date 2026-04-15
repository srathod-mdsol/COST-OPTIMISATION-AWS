"""
XLSX Formatter - Handles Excel XLSX export format
"""

import io
from typing import Any, List

from ..base_formatter import BaseFormatter, FormatType


class XLSXFormatter(BaseFormatter):
    """Formatter for Excel XLSX format"""
    
    def __init__(self, include_headers: bool = True, encoding: str = 'utf-8',
                 sheet_name: str = 'Data', freeze_panes: bool = True):
        """
        Initialize XLSX formatter
        
        Args:
            include_headers: Whether to include column headers
            encoding: Character encoding (for text conversion)
            sheet_name: Name of the Excel sheet
            freeze_panes: Whether to freeze header row
        """
        super().__init__(include_headers=include_headers, encoding=encoding)
        self.sheet_name = sheet_name
        self.freeze_panes = freeze_panes
    
    @property
    def format_type(self) -> FormatType:
        return FormatType.XLSX
    
    def format(self, data: Any, **kwargs) -> bytes:
        """
        Format data as XLSX bytes
        
        Args:
            data: Data to format (DataFrame, list of dicts, etc.)
            **kwargs: Additional options
            
        Returns:
            XLSX formatted bytes
        """
        if not self.validate_data(data):
            return b""
        
        # Try to use openpyxl for XLSX creation
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            self._errors.append("openpyxl not installed. Install with: pip install openpyxl")
            return b""
        
        # Convert data to records
        records = self._convert_to_records(data)
        
        if not records:
            self._warnings.append("No records to export")
            return b""
        
        # Create workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = self.sheet_name
        
        # Get headers from first record
        if records:
            headers = list(records[0].keys())
            
            # Write headers with styling
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="FF9900", end_color="FF9900", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center")
            
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
            
            # Freeze panes
            if self.freeze_panes:
                ws.freeze_panes = ws.cell(row=2, column=1)
            
            # Write data rows
            for row_idx, record in enumerate(records, 2):
                for col_idx, header in enumerate(headers, 1):
                    value = record.get(header, "")
                    cell_value = self._format_cell_value(value)
                    ws.cell(row=row_idx, column=col_idx, value=cell_value)
            
            # Auto-adjust column widths
            for col_idx, header in enumerate(headers, 1):
                max_length = len(str(header))
                for row_idx in range(2, len(records) + 2):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width
        
        # Save to bytes buffer
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()
    
    def _convert_to_records(self, data: Any) -> List[dict]:
        """Convert various data types to list of dictionaries"""
        # Handle pandas DataFrame
        try:
            import pandas as pd
            if isinstance(data, pd.DataFrame):
                # Convert datetime columns to strings
                for col in data.columns:
                    if pd.api.types.is_datetime64_any_dtype(data[col]):
                        data[col] = data[col].dt.isoformat()
                return data.to_dict('records')
        except ImportError:
            pass
        
        # Handle list of dictionaries
        if isinstance(data, list) and data and isinstance(data[0], dict):
            return data
        
        # Handle SQLAlchemy Row objects - convert to dict
        if isinstance(data, list) and data:
            first_item = data[0]
            # Check if it's a Row-like object
            if hasattr(first_item, '_mapping'):
                return [dict(row._mapping) for row in data]
            elif hasattr(first_item, 'keys'):
                return [dict(row) for row in data]
            elif hasattr(first_item, '__dict__'):
                # It's an object, convert attributes to dict
                return [{k: v for k, v in vars(row).items() if not k.startswith('_')} for row in data]
        
        # Handle single dictionary
        if isinstance(data, dict):
            return [data]
        
        # Handle other iterables
        try:
            return list(data)
        except:
            self._errors.append(f"Unsupported data type: {type(data)}")
            return []
    
    def _format_cell_value(self, value: Any) -> Any:
        """Format value for Excel cell"""
        if value is None:
            return ""
        if isinstance(value, (int, float, bool)):
            return value
        if isinstance(value, str):
            return value
        # Convert to string for other types
        return str(value)
    
    def format_to_bytes(self, data: Any, **kwargs) -> bytes:
        """Format data as XLSX bytes (alias for format)"""
        return self.format(data, **kwargs)
